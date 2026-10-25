import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_quotation():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    primary_color = RGBColor(27, 94, 32)      # Deep Emerald Green
    secondary_color = RGBColor(46, 125, 50)   # Green Accent
    dark_text = RGBColor(30, 41, 59)          # Slate 800
    gray_text = RGBColor(100, 116, 139)       # Slate 500

    # Header Top Bar
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_brand = p_top.add_run("BIENESTAR SIN EXCUSAS • TECNOLOGÍA NUTRICIONAL CLÍNICA\n")
    run_brand.font.name = "Calibri"
    run_brand.font.size = Pt(8.5)
    run_brand.font.bold = True
    run_brand.font.color.rgb = primary_color

    run_web = p_top.add_run("www.bienestarsinexcusas.site | +51 978 438 246")
    run_web.font.name = "Calibri"
    run_web.font.size = Pt(8.5)
    run_web.font.color.rgb = gray_text

    # Main Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("PROPUESTA DE SERVICIOS INTEGRALES Y MODERNIZACIÓN DIGITAL")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = primary_color

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Plataforma Web Clínica & Aplicativo Móvil Personalizado de Marca Blanca")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = secondary_color

    # Metadata Box Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_widths = [Inches(1.8), Inches(5.0)]
    for row in meta_table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w
            set_cell_margins(row.cells[i], top=100, bottom=100, left=140, right=140)

    fields = [
        ("DIRIGIDO A:", "Lic. Yajaira Alarcón Torres — Especialista en Nutrición (Puno, Perú)"),
        ("DE:", "Alberto Zegarra Ballón — Director General, Bienestar Sin Excusas"),
        ("FECHA DE EMISIÓN:", "24 de Octubre de 2026 | Vigencia: 15 días"),
        ("OBJETIVO:", "Implementación de Portal Web Oficial, App Móvil y Automatización de Planes de 28 Días")
    ]

    for idx, (label, val) in enumerate(fields):
        c0 = meta_table.cell(idx, 0)
        c1 = meta_table.cell(idx, 1)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = primary_color

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = dark_text
        if idx == 0:
            r1.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 1: Resumen Ejecutivo
    p_sec1 = doc.add_paragraph()
    r_sec1 = p_sec1.add_run("1. RESUMEN EJECUTIVO Y PROPUESTA DE VALOR")
    r_sec1.font.name = "Calibri"
    r_sec1.font.size = Pt(12)
    r_sec1.font.bold = True
    r_sec1.font.color.rgb = primary_color

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.15
    p_desc.paragraph_format.space_after = Pt(12)
    r_desc = p_desc.add_run(
        "La presente propuesta tiene como propósito equipar a la Lic. Yajaira Alarcón Torres con una infraestructura tecnológica "
        "de última generación desde el inicio de su ejercicio profesional en Puno y a nivel nacional. El sistema sustituye por completo "
        "el formato tradicional de archivos PDF y hojas de cálculo manuales por una suite digital interactiva bajo su propio nombre y marca, "
        "garantizando que sus pacientes vivan una experiencia moderna, organizada y de alto valor percibido."
    )
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(10)
    r_desc.font.color.rgb = dark_text

    # Section 2: Módulos del Sistema
    p_sec2 = doc.add_paragraph()
    r_sec2 = p_sec2.add_run("2. ARQUITECTURA DE LA PLATAFORMA (MÓDULOS INCLUIDOS)")
    r_sec2.font.name = "Calibri"
    r_sec2.font.size = Pt(12)
    r_sec2.font.bold = True
    r_sec2.font.color.rgb = primary_color

    modules = [
        ("Módulo 1: Web Oficial y Portal Clínico Personalizado", 
         "Página web moderna con el nombre y fotografía de Yajaira Alarcón, donde sus pacientes pueden conocer sus servicios, agendar citas, enviar datos clínicos y acceder a su perfil."),
        ("Módulo 2: Aplicativo Móvil PWA de Marca Blanca (iOS y Android)", 
         "App móvil instalable con 1 solo clic directamente en la pantalla de inicio del celular del paciente, con el logo, nombre y colores institucionales de su consultorio, sin pasar por tiendas ni descargas complejas."),
        ("Módulo 3: Automatización de Planes de 28 Días & Recetario", 
         "Generación de planes clínicos de 4 semanas (desayunos, almuerzos, cenas y colaciones) con cálculo nutricional, preparación detallada de recetas e integración de la lista de compras del mercado por semana."),
        ("Módulo 4: Asistente Virtual 24/7 y Canal Directo de Pacientes", 
         "Soporte interactivo para que el paciente resuelva dudas frecuentes sobre sus comidas, equivalencias de alimentos y recomendaciones de hidratación, manteniendo el vínculo activo con su especialista.")
    ]

    for mod_title, mod_body in modules:
        p_m = doc.add_paragraph()
        p_m.paragraph_format.space_after = Pt(4)
        p_m.paragraph_format.left_indent = Inches(0.2)
        r_mt = p_m.add_run(f"• {mod_title}: ")
        r_mt.font.name = "Calibri"
        r_mt.font.size = Pt(10)
        r_mt.font.bold = True
        r_mt.font.color.rgb = secondary_color

        r_mb = p_m.add_run(mod_body)
        r_mb.font.name = "Calibri"
        r_mb.font.size = Pt(9.5)
        r_mb.font.color.rgb = dark_text

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 3: Planes y Tarifas
    p_sec3 = doc.add_paragraph()
    r_sec3 = p_sec3.add_run("3. PLANES DE INVERSIÓN MENSUAL (EN SOLES PERUANOS - PEN)")
    r_sec3.font.name = "Calibri"
    r_sec3.font.size = Pt(12)
    r_sec3.font.bold = True
    r_sec3.font.color.rgb = primary_color

    plan_table = doc.add_table(rows=4, cols=4)
    plan_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    plan_table.autofit = False

    t_widths = [Inches(1.8), Inches(1.5), Inches(1.8), Inches(1.7)]
    headers = ["PLAN", "CAPACIDAD", "COSTO X PACIENTE", "INVERSIÓN MENSUAL"]
    
    for i, h in enumerate(headers):
        cell = plan_table.cell(0, i)
        cell.width = t_widths[i]
        set_cell_background(cell, "1B5E20")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    data_plans = [
        ("Plan 30\n(Recomendado Inicio)", "Hasta 30 pacientes", "S/ 13.33 / paciente", "S/ 400.00 / mes"),
        ("Plan 80", "Hasta 80 pacientes", "S/ 8.75 / paciente", "S/ 700.00 / mes"),
        ("Plan 200", "Hasta 200 pacientes", "S/ 6.00 / paciente", "S/ 1,200.00 / mes")
    ]

    for row_idx, data in enumerate(data_plans, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = plan_table.cell(row_idx, col_idx)
            cell.width = t_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = dark_text
            if col_idx in (0, 3):
                r.font.bold = True
            if row_idx == 1 and col_idx == 3:
                r.font.color.rgb = primary_color

    # Note about setup fee
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(6)
    p_note.paragraph_format.space_after = Pt(12)
    r_nb = p_note.add_run("★ BENEFICIO ESPECIAL DE LANZAMIENTO: ")
    r_nb.font.name = "Calibri"
    r_nb.font.size = Pt(9.5)
    r_nb.font.bold = True
    r_nb.font.color.rgb = primary_color
    r_nt = p_note.add_run("El costo de diseño, configuración técnica y desarrollo de la App y Portal Web (valorizado regularmente en S/ 800.00) está 100% BONIFICADO (GRATIS) desde el Plan 30 en adelante.")
    r_nt.font.name = "Calibri"
    r_nt.font.size = Pt(9.5)
    r_nt.font.color.rgb = dark_text

    # Section 4: Bono de Acción Rápida (48h)
    p_sec4 = doc.add_paragraph()
    r_sec4 = p_sec4.add_run("4. BONO DE ACCIÓN RÁPIDA (VÁLIDO POR 48 HORAS)")
    r_sec4.font.name = "Calibri"
    r_sec4.font.size = Pt(12)
    r_sec4.font.bold = True
    r_sec4.font.color.rgb = primary_color

    bono_table = doc.add_table(rows=1, cols=1)
    bono_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    bono_cell = bono_table.cell(0, 0)
    bono_cell.width = Inches(6.8)
    set_cell_background(bono_cell, "ECFDF5")
    set_cell_margins(bono_cell, top=140, bottom=140, left=180, right=180)

    p_b = bono_cell.paragraphs[0]
    p_b.paragraph_format.space_after = Pt(4)
    r_bt = p_b.add_run("🎁 BONIFICACIÓN ESPECIAL DE ARRANQUE: +15 DÍAS ADICIONALES SIN COSTO\n")
    r_bt.font.name = "Calibri"
    r_bt.font.size = Pt(10)
    r_bt.font.bold = True
    r_bt.font.color.rgb = primary_color

    r_bb = p_b.add_run(
        "Al confirmar el inicio del Plan 30 dentro de las próximas 48 horas (antes del lunes 26 de octubre a las 06:00 p. m.), "
        "se le otorgarán 45 DÍAS CONTINUOS de plataforma y aplicativo al precio regular de 1 mes (Ahorro directo de S/ 200.00), "
        "permitiéndole atender a sus primeros pacientes con amplio margen de rentabilidad."
    )
    r_bb.font.name = "Calibri"
    r_bb.font.size = Pt(9.5)
    r_bb.font.color.rgb = dark_text

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 5: Condiciones Comerciales y Cuentas de Abono
    p_sec5 = doc.add_paragraph()
    r_sec5 = p_sec5.add_run("5. CONDICIONES COMERCIALES Y MODALIDAD DE PAGO")
    r_sec5.font.name = "Calibri"
    r_sec5.font.size = Pt(12)
    r_sec5.font.bold = True
    r_sec5.font.color.rgb = primary_color

    p_cond = doc.add_paragraph()
    p_cond.paragraph_format.space_after = Pt(6)
    r_c1 = p_cond.add_run("• Facilidad 50% / 50%: ")
    r_c1.font.bold = True
    r_c1.font.color.rgb = secondary_color
    p_cond.add_run("Para iniciar la personalización y montaje de su App con su nombre y logo, se requiere un anticipo del 50% (S/ 200.00 para el Plan 30). El 50% restante (S/ 200.00) se cancela contra entrega formal de la plataforma activa.")
    
    p_cond2 = doc.add_paragraph()
    p_cond2.paragraph_format.space_after = Pt(10)
    r_c2 = p_cond2.add_run("• Tiempo de Entrega: ")
    r_c2.font.bold = True
    r_c2.font.color.rgb = secondary_color
    p_cond2.add_run("Aproximadamente 48 a 72 horas hábiles tras recibir sus datos y logotipo/foto profesional.")

    # Payment accounts box
    pay_table = doc.add_table(rows=1, cols=1)
    pay_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    pay_cell = pay_table.cell(0, 0)
    pay_cell.width = Inches(6.8)
    set_cell_background(pay_cell, "F8FAFC")
    set_cell_margins(pay_cell, top=140, bottom=140, left=180, right=180)

    p_pay = pay_cell.paragraphs[0]
    p_pay.paragraph_format.space_after = Pt(0)
    r_pt = p_pay.add_run("CUENTAS OFICIALES DE ABONO (PERÚ - SOLES):\n")
    r_pt.font.name = "Calibri"
    r_pt.font.size = Pt(9.5)
    r_pt.font.bold = True
    r_pt.font.color.rgb = primary_color

    r_pdet = p_pay.add_run(
        "• BBVA Banco Continental:\n"
        "  - Cuenta de Ahorros Soles: 0011-0131-0200371618\n"
        "  - Código Interbancario (CCI): 0111-31000-20037161806\n"
        "  - Titular: Alberto Gustavo Zegarra Ballón Quezada\n"
        "• Billetera Digital (YAPE / PLIN):\n"
        "  - Número directo: 978 438 246 (Alberto Zegarra)"
    )
    r_pdet.font.name = "Calibri"
    r_pdet.font.size = Pt(9)
    r_pdet.font.color.rgb = dark_text

    # Signature
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(20)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sig1 = p_sig.add_run("Alberto Zegarra Ballón\n")
    r_sig1.font.name = "Calibri"
    r_sig1.font.size = Pt(11)
    r_sig1.font.bold = True
    r_sig1.font.color.rgb = primary_color

    r_sig2 = p_sig.add_run("Director General | Bienestar Sin Excusas\nTecnología Nutricional Clínica\nWhatsApp: +51 978 438 246")
    r_sig2.font.name = "Calibri"
    r_sig2.font.size = Pt(9)
    r_sig2.font.color.rgb = gray_text

    output_path = r"C:\Users\Alberto\Downloads\Cotizacion_Yajaira_Alarcon_Torres.docx"
    doc.save(output_path)
    print(f"Quotation successfully generated at: {output_path}")

    # Also save a copy in repo
    repo_copy = r"C:\Users\Alberto\Downloads\CRM\Cotizacion_Yajaira_Alarcon_Torres.docx"
    doc.save(repo_copy)
    print(f"Repo copy saved at: {repo_copy}")

if __name__ == "__main__":
    create_quotation()
